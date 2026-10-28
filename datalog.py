"""
Registro de datos de cada consulta a la tienda (una fila por consulta HTTP).
- Se guarda en DATA_DIR/consultas.csv (persistente).
- Se puede descargar como Excel (.xlsx) o CSV desde el panel.
- Opcional: se copia a Google Sheets mediante un Web App de Apps Script (SHEETS_WEBHOOK_URL).
"""
import asyncio
import csv
import io
import time
from datetime import datetime
from pathlib import Path

COLUMNS = [
    ("fecha_hora", "Fecha y hora (Bogotá)"),
    ("fecha", "Fecha"),
    ("hora", "Hora del día"),
    ("dia_semana", "Día de la semana"),
    ("revision", "N° revisión"),
    ("motivo", "Motivo"),
    ("modo", "Modo"),
    ("fuente", "Fuente"),
    ("termino", "Término / producto"),
    ("http", "Código HTTP"),
    ("ok", "Éxito (1/0)"),
    ("ms", "Tiempo de respuesta (ms)"),
    ("kb", "Tamaño respuesta (KB)"),
    ("productos", "Productos devueltos"),
    ("coincidencias", "Coincidencias calendario"),
    ("con_stock", "Calendario con stock (1/0)"),
    ("unidades", "Unidades disponibles"),
    ("precio", "Precio (COP)"),
    ("seg_desde_anterior", "Segundos desde consulta anterior"),
    ("error", "Error"),
]
KEYS = [k for k, _ in COLUMNS]
HEADERS = [h for _, h in COLUMNS]
DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]


class QueryLog:
    def __init__(self, path: Path, webhook: str = "", token: str = ""):
        self.path = path
        self.webhook = webhook.strip()
        self.token = token.strip()
        self.pending: list[list] = []        # filas aún no enviadas a Sheets
        self.last_ts: float | None = None
        self.total = self._count()
        self.sheets_last_ok: str | None = None
        self.sheets_last_error: str | None = None
        self.sheets_sent = 0
        self.flush_event = asyncio.Event()
        if not self.path.exists() or self.path.stat().st_size == 0:
            with self.path.open("w", newline="", encoding="utf-8") as f:
                csv.writer(f).writerow(KEYS)

    def _count(self) -> int:
        try:
            with self.path.open(encoding="utf-8") as f:
                return max(0, sum(1 for _ in f) - 1)
        except FileNotFoundError:
            return 0

    # ------------------------------------------------------------ registrar
    def record(self, when: datetime, **fields) -> dict:
        ts = when.timestamp()
        row = {
            "fecha_hora": when.strftime("%Y-%m-%d %H:%M:%S"),
            "fecha": when.strftime("%Y-%m-%d"),
            "hora": when.hour,
            "dia_semana": DIAS[when.weekday()],
            "seg_desde_anterior": round(ts - self.last_ts, 1) if self.last_ts else "",
            **fields,
        }
        self.last_ts = ts
        values = [row.get(k, "") for k in KEYS]
        with self.path.open("a", newline="", encoding="utf-8") as f:
            csv.writer(f).writerow(values)
        self.total += 1
        if self.webhook:
            self.pending.append(values)
            if len(self.pending) > 20000:      # nunca crece sin límite si Sheets falla días
                self.pending = self.pending[-20000:]
            if row.get("coincidencias"):       # si aparece el calendario, sube ya
                self.flush_event.set()
        return row

    def rows(self, limit: int | None = None) -> list[list[str]]:
        with self.path.open(encoding="utf-8") as f:
            data = list(csv.reader(f))[1:]
        return data[-limit:] if limit else data

    def clear(self):
        with self.path.open("w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerow(KEYS)
        self.total = 0
        self.pending.clear()

    # ------------------------------------------------------------ exportar
    def to_csv(self) -> bytes:
        out = io.StringIO()
        w = csv.writer(out)
        w.writerow(HEADERS)
        w.writerows(self.rows())
        return ("﻿" + out.getvalue()).encode("utf-8")   # BOM: tildes bien en Excel

    def to_xlsx(self) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo

        wb = Workbook()
        ws = wb.active
        ws.title = "Consultas"
        ws.append(HEADERS)
        numeric = {"hora", "revision", "http", "ok", "ms", "kb", "productos", "coincidencias",
                   "con_stock", "unidades", "precio", "seg_desde_anterior"}
        rows = self.rows()
        for r in rows:
            out = []
            for k, v in zip(KEYS, r):
                if k == "fecha_hora" and v:
                    out.append(datetime.strptime(v, "%Y-%m-%d %H:%M:%S"))
                elif k == "fecha" and v:
                    out.append(datetime.strptime(v, "%Y-%m-%d").date())
                elif k in numeric and v != "":
                    try:
                        out.append(float(v) if "." in v else int(v))
                    except ValueError:
                        out.append(v)
                else:
                    out.append(v)
            ws.append(out)
        head = Font(bold=True, color="FFFFFF")
        fill = PatternFill("solid", fgColor="D6336C")
        for c in ws[1]:
            c.font, c.fill = head, fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "A2"
        widths = {"fecha_hora": 20, "fecha": 12, "fuente": 18, "termino": 22, "error": 40, "dia_semana": 12}
        for i, k in enumerate(KEYS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(k, 13)
        for cell in ws["A"][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm:ss"
        for cell in ws["B"][1:]:
            cell.number_format = "yyyy-mm-dd"
        if rows:
            ref = f"A1:{get_column_letter(len(KEYS))}{len(rows) + 1}"
            t = Table(displayName="Consultas", ref=ref)
            t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium3", showRowStripes=True)
            ws.add_table(t)

        # Hoja de diccionario: qué significa cada columna (útil para el análisis)
        d = wb.create_sheet("Diccionario")
        d.append(["Columna", "Descripción"])
        for (k, h), desc in zip(COLUMNS, DESCRIPCIONES):
            d.append([h, desc])
        for c in d[1]:
            c.font, c.fill = head, fill
        d.column_dimensions["A"].width = 32
        d.column_dimensions["B"].width = 90

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ------------------------------------------------------------ Google Sheets
    def sheets_status(self) -> dict:
        return {
            "enabled": bool(self.webhook),
            "pending": len(self.pending),
            "sent": self.sheets_sent,
            "last_ok": self.sheets_last_ok,
            "last_error": self.sheets_last_error,
        }

    async def flush(self, client, log) -> bool:
        if not self.webhook or not self.pending:
            return True
        batch = self.pending[:500]
        try:
            r = await client.post(self.webhook, json={"token": self.token, "headers": HEADERS, "rows": batch},
                                  timeout=30)
            data = r.json()
            if not data.get("ok"):
                raise RuntimeError(data.get("error") or f"HTTP {r.status_code}")
        except Exception as e:
            msg = str(e) or e.__class__.__name__
            if msg != self.sheets_last_error:
                log("warn", f"Google Sheets no recibió los datos ({msg}); se reintenta luego")
            self.sheets_last_error = msg
            return False
        del self.pending[:len(batch)]
        self.sheets_sent += len(batch)
        self.sheets_last_ok = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S")
        if self.sheets_last_error:
            log("info", "Google Sheets volvió a recibir datos ✔")
        self.sheets_last_error = None
        log("debug", f"{len(batch)} filas enviadas a Google Sheets")
        if self.pending:
            return await self.flush(client, log)
        return True

    async def sync_forever(self, client, log, every_seconds: int = 300):
        if not self.webhook:
            return
        log("info", "Registro de consultas conectado a Google Sheets")
        while True:
            try:
                await asyncio.wait_for(self.flush_event.wait(), every_seconds)
            except asyncio.TimeoutError:
                pass
            self.flush_event.clear()
            try:
                await self.flush(client, log)
            except Exception as e:  # nunca debe tumbar el monitor
                log("warn", f"Sincronización con Sheets falló: {e!r}")


DESCRIPCIONES = [
    "Momento exacto de la consulta, hora de Bogotá (UTC-5).",
    "Solo la fecha, útil para agrupar por día.",
    "Hora del día (0–23), útil para ver patrones por franja horaria.",
    "Día de la semana de la consulta.",
    "Número de revisión del monitor. Una revisión manual puede incluir varias consultas.",
    "Por qué se hizo: programada, manual, inicio o reanudación.",
    "Modo del monitor: normal, ventana (10–20 nov) o turbo (calendario publicado sin stock).",
    "API de VTEX consultada: catálogo, intelligent-search o producto (consulta directa en turbo).",
    "Texto buscado o ID del producto consultado.",
    "Código de respuesta HTTP de la tienda (200/206 = bien; 403/429/5xx = problema).",
    "1 si la consulta respondió bien, 0 si falló.",
    "Milisegundos que tardó la tienda en responder.",
    "Tamaño de la respuesta en kilobytes.",
    "Cantidad de productos que devolvió la búsqueda.",
    "Productos cuyo nombre coincide con los filtros del calendario.",
    "1 si alguna coincidencia tenía stock en esa consulta.",
    "Suma de unidades disponibles reportadas para las coincidencias.",
    "Precio más bajo encontrado para las coincidencias, en pesos.",
    "Segundos transcurridos desde la consulta anterior (permite verificar el intervalo aleatorio).",
    "Mensaje de error si la consulta falló.",
]
